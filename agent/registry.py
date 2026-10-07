"""Workflow registry.

Excel is the source of truth for *what* a workflow is (name, trigger, inputs,
ordered steps, decision logic, tools, expected output). Python modules in
`workflows/` supply *how* each step runs. They are bound together here:

    @step("WF001")            # appended in definition order
    def load_inventory(ctx): ...

Adding workflow #11 = add a row to the Excel sheet + drop one file into
`workflows/`. Nothing else (router, engine, CLI, tests) changes.
"""
from __future__ import annotations

import importlib
import pkgutil
from dataclasses import dataclass, field
from typing import Callable

import openpyxl

from . import config
from .models import WorkflowSpec

StepFn = Callable[["RunContext"], "str | None"]  # noqa: F821


@dataclass
class StepHandler:
    fn: StepFn
    optional: bool = False


@dataclass
class WorkflowMeta:
    keywords: list[str] = field(default_factory=list)       # extra routing vocabulary
    param_keys: list[str] = field(default_factory=list)     # params the LLM router may fill
    extractor: Callable[[str], dict] | None = None          # regex param extraction from request


_HANDLERS: dict[str, list[StepHandler]] = {}
_META: dict[str, WorkflowMeta] = {}


def step(workflow_id: str, optional: bool = False):
    """Register the next step of a workflow (order of definition = order of execution).
    optional=True: if the step fails the run continues and the failure is reported."""
    def deco(fn: StepFn) -> StepFn:
        _HANDLERS.setdefault(workflow_id, []).append(StepHandler(fn, optional))
        return fn
    return deco


def workflow_meta(workflow_id: str, keywords=None, param_keys=None, extractor=None) -> None:
    _META[workflow_id] = WorkflowMeta(list(keywords or []), list(param_keys or []), extractor)


class Registry:
    def __init__(self):
        self.specs: dict[str, WorkflowSpec] = {}
        self.test_requests: list[dict] = []
        self.handlers = _HANDLERS
        self.meta = _META
        self.problems: list[str] = []

    # -- loading -------------------------------------------------------------
    @classmethod
    def load(cls, strict: bool = True) -> "Registry":
        from .tools import load_all
        load_all()
        import workflows  # noqa: F401  (package __init__ imports every module)
        reg = cls()
        reg._read_excel()
        reg._validate(strict)
        return reg

    def _read_excel(self) -> None:
        wb = openpyxl.load_workbook(config.WORKFLOW_XLSX, data_only=True)
        ws = wb["Workflows"]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip() for h in rows[0]]
        for r in rows[1:]:
            if not r or not r[0]:
                continue
            d = dict(zip(header, ["" if v is None else str(v).strip() for v in r]))
            steps = [s.strip() for s in d["Steps"].replace("->", "→").split("→") if s.strip()]
            self.specs[d["Workflow_ID"]] = WorkflowSpec(
                id=d["Workflow_ID"], name=d["Workflow_Name"], trigger=d["Trigger"],
                inputs=d["Inputs"], steps=steps, decision_logic=d["Decision_Logic"],
                tools=d["Tools_Required"], expected_output=d["Expected_Output"])
        if "Test_Questions" in wb.sheetnames:
            tws = wb["Test_Questions"]
            trows = list(tws.iter_rows(values_only=True))
            th = [str(h).strip() for h in trows[0]]
            for r in trows[1:]:
                if r and r[0]:
                    self.test_requests.append(dict(zip(th, ["" if v is None else str(v) for v in r])))

    def _validate(self, strict: bool) -> None:
        for wid, spec in self.specs.items():
            handlers = self.handlers.get(wid)
            if not handlers:
                self.problems.append(f"{wid} ({spec.name}) is in Excel but has no implementation in workflows/")
            elif len(handlers) != len(spec.steps):
                self.problems.append(
                    f"{wid}: Excel defines {len(spec.steps)} steps but code registers {len(handlers)}")
        for wid in self.handlers:
            if wid not in self.specs:
                self.problems.append(f"{wid} has code but no row in the Excel workflow sheet")
        mismatched = [p for p in self.problems if "steps but code" in p or "no row in the Excel" in p]
        if strict and mismatched:
            raise RuntimeError("Workflow registry inconsistent:\n  " + "\n  ".join(mismatched))

    # -- helpers -------------------------------------------------------------
    def implemented(self, wid: str) -> bool:
        return wid in self.handlers and len(self.handlers[wid]) == len(self.specs[wid].steps)
